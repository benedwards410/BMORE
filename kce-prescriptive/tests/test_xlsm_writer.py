import hashlib
import os
import tempfile
import unittest
import zipfile

import _path  # noqa: F401  (sets sys.path)
from openpyxl import Workbook, load_workbook

from prescriptive import xlsm_writer as w


def _book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Lighting Inventory"
    ws["A1"] = "label"
    ws["B2"] = 5
    ws["C2"] = "=B2*2"
    ws["D3"] = "=B2+1"
    other = wb.create_sheet("Other")
    other["A1"] = "untouched"
    wb.save(path)


class WriterTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.src = os.path.join(self.tmp, "t.xlsx")
        self.out = os.path.join(self.tmp, "o.xlsx")
        _book(self.src)

    def test_values_land_and_formulas_survive(self):
        w.write(self.src, self.out, {"Lighting Inventory": {"A1": "A & <B>", "B2": 7, "E2": 1.5, "A9": "new row", "Z2": None}},
                strip_sheets=["Lighting Inventory"])
        wb = load_workbook(self.out)
        ws = wb["Lighting Inventory"]
        self.assertEqual(ws["A1"].value, "A & <B>")
        self.assertEqual(ws["B2"].value, 7)
        self.assertEqual(ws["E2"].value, 1.5)
        self.assertEqual(ws["A9"].value, "new row")
        self.assertEqual(ws["C2"].value, "=B2*2")
        self.assertEqual(wb["Other"]["A1"].value, "untouched")

    def test_refuses_to_overwrite_formula(self):
        with self.assertRaises(w.FormulaCellError):
            w.write(self.src, self.out, {"Lighting Inventory": {"C2": 3}})

    def test_if_not_formula(self):
        w.write(self.src, self.out, {"Lighting Inventory": {"C2": w.IfNotFormula("x"), "C3": w.IfNotFormula("y")}})
        ws = load_workbook(self.out)["Lighting Inventory"]
        self.assertEqual(ws["C2"].value, "=B2*2")
        self.assertEqual(ws["C3"].value, "y")

    def test_untouched_parts_are_byte_identical(self):
        w.write(self.src, self.out, {"Lighting Inventory": {"A1": "x"}})
        a, b = zipfile.ZipFile(self.src), zipfile.ZipFile(self.out)
        self.assertEqual(a.namelist(), b.namelist())
        changed = [n for n in a.namelist() if hashlib.sha256(a.read(n)).digest() != hashlib.sha256(b.read(n)).digest()]
        self.assertTrue(set(changed) <= {"xl/workbook.xml", "xl/worksheets/sheet1.xml"}, changed)
        self.assertIn("xl/worksheets/sheet1.xml", changed)
        self.assertIn('fullCalcOnLoad="1"', b.read("xl/workbook.xml").decode())

    def test_strip_keeps_every_cell(self):
        xml = ('<row r="1"><c r="A1" s="1"/><c r="B1" s="2" t="str"><f>X()</f><v xml:space="preserve">old </v></c>'
               '<c r="C1" t="str"><f t="shared" si="1"/><v/></c><c r="D1" t="s"><v>3</v></c></row>')
        out = w.strip_cached_values(xml)
        self.assertEqual(out, '<row r="1"><c r="A1" s="1"/><c r="B1" s="2"><f>X()</f></c>'
                              '<c r="C1"><f t="shared" si="1"/></c><c r="D1" t="s"><v>3</v></c></row>')


if __name__ == "__main__":
    unittest.main()
