"""T&C fill, proposal numbers, audit sheet round trip, reference integrity."""

import os
import tempfile
import unittest

from helpers import CATALOG, CONFIG, TABLES, example_project, run
from openpyxl import load_workbook

from prescriptive import audit, proposal, terms_pdf
from prescriptive.build import TNC_TEMPLATE
from prescriptive.project import load_existing


class TermsTest(unittest.TestCase):
    def test_fill_and_read_back(self):
        project, *_ = run(example_project())
        out = os.path.join(tempfile.mkdtemp(), "tc.pdf")
        terms_pdf.fill(TNC_TEMPLATE, out, project, CONFIG)
        vals, payee = terms_pdf.read_back(out)
        self.assertEqual(vals["Project Name"], "Example Holdings LLC")
        self.assertEqual(vals["BGE Electric Account Number"], "1234567890")
        self.assertEqual(vals["Street Address where the work is taking place"], "100 Example Ave, Suite 2")
        self.assertEqual(vals["Street Address"], "5548 April Journey")
        self.assertEqual(payee, "contractor")

    def test_customer_payee_leaves_assignee_blank(self):
        project, *_ = run(example_project(payee="customer"))
        out = os.path.join(tempfile.mkdtemp(), "tc.pdf")
        terms_pdf.fill(TNC_TEMPLATE, out, project, CONFIG)
        vals, payee = terms_pdf.read_back(out)
        self.assertEqual(payee, "customer")
        self.assertEqual(vals["Street Address"], "")


class ProposalTest(unittest.TestCase):
    def test_quote_arithmetic(self):
        # 82 kits at $100 + 1% recycling + 3% tax; 246 lamps / 82 ballasts at 60 h/wk with the config's maintenance assumptions
        p = example_project()
        p["lines"] = [{"area": "Office", "int_ext": "Interior - Open", "space_type": "Office - Open Plan", "over_300": "Yes",
                       "existing": "2x4-3L-F32T8", "existing_qty": 82, "product": "VEKT2X4-8T", "control": "IFS108",
                       "daylight_qty": 82, "unit_price": 100}]
        project, resolved, _, est, _ = run(p)
        n = proposal.numbers(project, resolved, est, CONFIG, incentive=1000)
        self.assertEqual((n["subtotal"], n["recycling"], n["tax"], n["total"], n["net"]), (8200, 82.0, 246.0, 8528.0, 7528.0))
        self.assertEqual((n["lamps"], n["ballasts"]), (246, 82))
        self.assertAlmostEqual(n["maint_usd"], 996.71, places=2)
        self.assertAlmostEqual(n["saved_kwh"], (82 * 88 - 82 * 25) * 3120 / 1000, places=6)

    def test_docx_builds(self):
        project, resolved, _, est, _ = run(example_project())
        out = os.path.join(tempfile.mkdtemp(), "p.docx")
        proposal.build(out, project, resolved, est, TABLES, CONFIG)
        self.assertTrue(os.path.getsize(out) > 5000)


class AuditSheetTest(unittest.TestCase):
    def test_round_trip(self):
        d = tempfile.mkdtemp()
        path = audit.write_template(os.path.join(d, "a.xlsx"))
        wb = load_workbook(path)
        pj, au = wb["Project"], wb["Audit"]
        values = {"customer_of_record": "Example Holdings LLC", "site_address": "100 Example Ave", "city": "Columbia",
                  "zip": "21045", "bge_account": "1234567890", "building_type": "Office", "heating_fuel": "Gas",
                  "contact.name": "Pat Example", "contact.email": "pat@example.com"}
        for row in pj.iter_rows(min_row=2):
            if row[2].value in values:
                row[1].value = values[row[2].value]
        keys = [c.value for c in au[2]]
        line = {"area": "Open office", "int_ext": "Interior - Open", "space_type": "Office - Open Plan", "over_300": "Yes",
                "existing": "2x4-3L-F32T8 · 2x4, 3-lamp F32T8", "existing_qty": 10,
                "product": "VEKT2X4-8T · 25W 2x4 center-basket kit", "control": "IFS108 · Occupancy + daylight", "daylight_qty": 4}
        for k, v in line.items():
            au.cell(row=3, column=keys.index(k) + 1, value=v)
        wb.save(path)
        project, notes = audit.to_project(path)
        self.assertEqual(notes, [])
        self.assertEqual(project["lines"][0]["product"], "VEKT2X4-8T")
        self.assertEqual(project["audit_totals"]["existing_qty"], 10)
        _, _, _, est, _ = run(project)
        self.assertEqual(est["total"], 600 + 200)


class ReferenceIntegrityTest(unittest.TestCase):
    def test_existing_fixtures_match_template(self):
        for k, ex in load_existing()["fixtures"].items():
            self.assertIn(ex["code"], TABLES["wattage"], k)
            self.assertIn(ex["code"], TABLES["device_categories"][ex["category"]], k)

    def test_catalog_measures_and_controls_exist(self):
        for k, fx in CATALOG["fixtures"].items():
            if fx.get("measure"):
                self.assertIn(fx["measure"], TABLES["measures"], k)
        for k, c in CATALOG["controls"].items():
            self.assertIn(c["control_type"], TABLES["control_types"], k)
            self.assertIn(c["control_select"], TABLES["control_selects"][c["control_type"]], k)
            if c.get("nlc"):
                self.assertIn(c["nlc"]["secondary"], ["Daylight Harvesting", "Trim", "Zoning", "Peak Demand Control", "Other"])


if __name__ == "__main__":
    unittest.main()
