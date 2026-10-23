"""Customer proposal (.docx), laid out like KCE's 'BGE Lighting Retrofit Quote' (Oct 2026 format).

Sections: summary, scope by area, pricing, estimated BGE incentive, savings and
payback, how it works, acceptance. Dollars appear here by decision (price -
BGE incentive = net); the file stays a draft until Ben approves it.
"""

import datetime as dt
import os
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .workbook import street


def money(x, cents=True):
    s = f"${abs(x):,.2f}" if cents else f"${abs(x):,.0f}"
    return "-" + s if x < 0 else s


def _date(project):
    d = project.get("date")
    d = dt.date.fromisoformat(d) if d else dt.date.today()
    return f"{d:%b} {d.day}, {d.year}"


def numbers(project, resolved, est, config, incentive=None):
    """Every figure the proposal prints, in one place (also used by the portal sheet and tests)."""
    s = project["savings"]
    mt = s["maintenance"]
    hours = s["hours_per_week"] * 52
    subtotal = sum((r.get("unit_price") or 0) * (r.get("qty") or 0) for r in resolved)
    recycling = round(subtotal * project["pricing"]["recycling_pct"] / 100, 2)
    tax = round(subtotal * project["pricing"]["tax_pct"] / 100, 2)
    lift = project["pricing"].get("lift_fee") or 0
    total = round(subtotal + recycling + tax + lift, 2)
    incentive = est["total"] if incentive is None else incentive
    net = round(total - incentive, 2)

    ex_w = sum((r.get("existing_watts") or 0) * (r.get("existing_qty") or 0) for r in resolved)
    new_w = sum(((r.get("fixture") or {}).get("field_watts") or (r.get("fixture") or {}).get("watts") or 0) * (r.get("qty") or 0)
                for r in resolved)
    ex_kwh = ex_w * hours / 1000
    new_kwh = new_w * hours / 1000
    saved_kwh = ex_kwh - new_kwh
    energy_usd = saved_kwh * s["rate_cents_per_kwh"] / 100
    lamps = sum((r.get("existing_lamps") or 0) * (r.get("existing_qty") or 0) for r in resolved if r.get("existing_fluorescent"))
    ballasts = sum((r.get("existing_ballasts") or 0) * (r.get("existing_qty") or 0) for r in resolved if r.get("existing_fluorescent"))
    lamp_cost = mt["lamp_minutes"] / 60 * mt["labor_rate"] + mt["lamp_part"]
    ballast_cost = mt["ballast_minutes"] / 60 * mt["labor_rate"] + mt["ballast_part"]
    maint_usd = lamps * hours / mt["lamp_life_hours"] * lamp_cost + ballasts * hours / mt["ballast_life_hours"] * ballast_cost
    annual = energy_usd + maint_usd
    return {
        "subtotal": subtotal, "recycling": recycling, "tax": tax, "lift": lift, "total": total,
        "incentive": incentive, "net": net, "hours": hours,
        "ex_kwh": ex_kwh, "new_kwh": new_kwh, "saved_kwh": saved_kwh,
        "pct_saved": saved_kwh / ex_kwh * 100 if ex_kwh else 0,
        "energy_usd": energy_usd, "maint_usd": maint_usd, "lamps": lamps, "ballasts": ballasts,
        "annual_usd": annual, "payback_years": net / annual if annual else None,
        "fixtures": sum(r.get("qty") or 0 for r in resolved),
        "sensors": sum(r.get("control_qty") or 0 for r in resolved
                       if r.get("control_product") and r.get("control_type") != "Other Non-Incentivized Option"),
    }


def _table(doc, header, rows, bold_last=0):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, header):
        c.text = ""
        c.paragraphs[0].add_run(str(h)).bold = True
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = "" if v is None else str(v)
            if bold_last and i >= len(rows) - bold_last:
                for run in c.paragraphs[0].runs:
                    run.bold = True
    return t


def _para(doc, parts, style=None):
    p = doc.add_paragraph(style=style)
    for text, bold in parts:
        p.add_run(text).bold = bold
    return p


def _measure_name(code, tables):
    m = tables["measures"].get(code) or tables["controls"].get(code) or {}
    name = (m.get("summary_description") or m.get("select") or code).strip()
    name = name.replace(" fixtures", "").replace("LED integrated retrofit kit for ", "LED integrated retrofit kit, ")
    return name.replace("Dual daylight & occupancy sensor (DOS)", "Dual daylight and occupancy sensor")


def build(out_docx, project, resolved, est, tables, config, incentive=None, official=False):
    n = numbers(project, resolved, est, config, incentive)
    co = config["company"]
    p = project["project"]
    doc = Document()
    for st in ("Normal",):
        doc.styles[st].font.size = Pt(10.5)

    doc.add_heading(f"BGE Lighting Retrofit Quote – {n['fixtures']} Fixtures", level=1)
    by = " · ".join(x for x in (co["name"], _date(project), project.get("prepared_by")) if x)
    doc.add_paragraph(by)

    # Summary
    doc.add_heading("Project summary", level=2)
    est_word = "BGE incentive" if official else "estimated BGE incentive"
    _para(doc, [(f"{co['name']} will retrofit {n['fixtures']} fixtures to LED for a total project cost of ", False),
                (money(n["total"]), True), (f". The {est_word} is ", False), (money(n["incentive"]), True),
                (", bringing the net investment to ", False), (money(n["net"]), True), (".", False)])

    by_product = {}
    for r in resolved:
        fx = r.get("fixture") or {}
        by_product[fx.get("label") or r.get("product")] = by_product.get(fx.get("label") or r.get("product"), 0) + (r.get("qty") or 0)
    fixtures_txt = f"{n['fixtures']} LED fixtures (" + ", ".join(f"{q} {lbl}" for lbl, q in by_product.items()) + ")"
    ctrl_labels = sorted({(r.get("control_product") or {}).get("label", "") for r in resolved if r.get("control_product")})
    if n["sensors"]:
        sensors_txt = f"{n['sensors']} {' / '.join(l.lower() for l in ctrl_labels)} sensors" + (", one per fixture" if n["sensors"] == n["fixtures"] else "")
    else:
        sensors_txt = "None"
    site = f"{street(project)} {p.get('city', '')}, {p.get('state', 'MD')} {p.get('zip', '')}".strip()
    _table(doc, ["Item", "Detail"], [
        ["Customer", p.get("customer_of_record") or ""],
        ["Site address", site],
        ["BGE account number", p.get("bge_account") or ""],
        ["Program", "BGE Energy Solutions for Business, prescriptive retrofit lighting"],
        ["Fixtures", fixtures_txt],
        ["Sensors", sensors_txt],
        ["Energy saved", f"{n['saved_kwh']:,.0f} kWh per year, {n['pct_saved']:.0f}% of existing lighting energy"],
    ])

    # Scope
    doc.add_heading("Scope by area", level=2)
    all_kits = all(r.get("measure_type") == "LED RETROFIT KITS" for r in resolved)
    all_ctrl = n["sensors"] == n["fixtures"] and n["fixtures"] > 0
    intro = ("Every existing fixture gets an LED retrofit kit inside its current housing"
             if all_kits else "Each existing fixture is replaced one-for-one with the LED product shown")
    if all_ctrl and len(ctrl_labels) == 1:
        intro += f", and every fixture gets an {ctrl_labels[0].lower()} sensor"
    doc.add_paragraph(intro + ".")
    rows = []
    for r in resolved:
        fx = r.get("fixture") or {}
        cp = r.get("control_product") or {}
        rows.append([r.get("area"), f"{r.get('existing_label')} ({r.get('existing_watts')}W)",
                     fx.get("label") or r.get("product"), cp.get("label") or "None", r.get("qty")])
    rows.append(["Total", "", "", f"{n['sensors']} sensors" if n["sensors"] else "", n["fixtures"]])
    _table(doc, ["Area", "Existing fixture", "New LED", "Sensor", "Qty"], rows, bold_last=1)

    # Pricing
    doc.add_heading("Pricing", level=2)
    prices = {}
    for r in resolved:
        fx = r.get("fixture") or {}
        prices.setdefault(r.get("product"), (fx.get("short") or fx.get("label") or r.get("product"), r.get("unit_price") or 0))
    uniq = {pr for _, pr in prices.values()}
    shorts = [s for s, _ in prices.values()]
    each = (f"Each {' or '.join(s.replace(' kits', '') for s in shorts)} kit with sensor is {money(uniq.pop())}."
            if len(uniq) == 1 and all_kits else "Unit prices include the sensor where one is listed.")
    lift_txt = " Includes lift rental." if n["lift"] else " No lift is needed."
    doc.add_paragraph(f"The installed price is {money(n['total'])}, including recycling and tax. {each}{lift_txt}")
    tail = [["Lamp and ballast recycling (%g%%)" % project["pricing"]["recycling_pct"], money(n["recycling"])],
            ["Tax (%g%%)" % project["pricing"]["tax_pct"], money(n["tax"])]]
    if n["lift"]:
        tail.append(["Lift rental", money(n["lift"])])
    tail += [["Total project cost", money(n["total"])],
             [("BGE incentive" if official else "Estimated BGE incentive"), money(-n["incentive"])],
             ["Net investment", money(n["net"])]]
    if len(prices) <= 3:
        pids = list(prices)
        areas = []
        grid = {}
        for r in resolved:
            a = r.get("area")
            if a not in grid:
                areas.append(a)
                grid[a] = {}
            grid[a][r["product"]] = grid[a].get(r["product"], 0) + (r.get("qty") or 0)
        header = ["Area"] + [f"{prices[pid][0]} @ {money(prices[pid][1])}" for pid in pids] + ["Line total"]
        body = []
        for a in areas:
            line_total = sum(grid[a].get(pid, 0) * prices[pid][1] for pid in pids)
            body.append([a] + [grid[a].get(pid) or "" for pid in pids] + [money(line_total)])
        body.append(["Subtotal"] + [sum(grid[a].get(pid, 0) for a in areas) for pid in pids] + [money(n["subtotal"])])
        pad = [""] * len(pids)
        body += [[label] + pad + [val] for label, val in tail]
    else:
        header = ["Item", "Qty", "Unit price", "Line total"]
        body = []
        for pid, (short, pr) in prices.items():
            q = sum(r.get("qty") or 0 for r in resolved if r.get("product") == pid)
            body.append([short, q, money(pr), money(q * pr)])
        body.append(["Subtotal", n["fixtures"], "", money(n["subtotal"])])
        body += [[label, "", "", val] for label, val in tail]
    _table(doc, header, body, bold_last=3)

    # Incentive
    doc.add_heading("BGE incentive" if official else "Estimated BGE incentive", level=2)
    payee = project.get("payee", "contractor")
    paid = (f", assigned to and paid directly to {co['name']}" if payee != "customer" else ", paid by BGE to you")
    all_pass = all(not e["reasons"] for e in est["lines"])
    verb = "pay" if official else "are estimated to pay"
    doc.add_paragraph(f"BGE's prescriptive lighting measures {verb} {money(n['incentive'])}{paid}."
                      + (" Every fixture clears its code's minimum watts reduced." if all_pass else ""))
    rows = []
    ordered = sorted(est["by_measure"].items(), key=lambda kv: kv[0].startswith("LC"))   # lighting first, then controls
    for code, b in ordered:
        name = _measure_name(code, tables)
        if code.startswith("RL"):
            reds = sorted({e.get("watts_reduced_per_unit") for e in est["lines"] if e["measure"] == code and e.get("watts_reduced_per_unit") is not None})
            mn = tables["measures"][code]["min_watts_reduced"]
            unit = tables["measures"][code]["unit"].lower()
            rng = f"{reds[0]}W" if len(reds) == 1 else f"{reds[0]}–{reds[-1]}W"
            detail = f"{name} ({rng} reduced{' per fixture' if code in ('RL 111', 'RL 111A', 'RL 111B', 'RL 121', 'RL 167') else ''}; {mn}W min.)"
            rate = f"{money(b['rate'], cents=False)} / {unit}"
        else:
            ws = sorted({(r.get('fixture') or {}).get('watts') for r in resolved if r.get('control_code') == code and (r.get('fixture') or {}).get('watts')})
            rng = (f"{ws[0]:g}W" if len(ws) == 1 else f"{ws[0]:g}–{ws[-1]:g}W") if ws else "?"
            detail = f"{name} ({rng} controlled; {tables['controls'][code]['min_watts']}W min.)"
            rate = f"{money(b['rate'], cents=False)} / fixture"
        rows.append([code, detail, b["units"], rate, money(b["amount"])])
    rows.append(["Total", "", "", "", money(est["total"])])
    _table(doc, ["BGE code", "Measure", "Qty", "Rate", "Incentive"], rows, bold_last=1)
    doc.add_paragraph(f"Rates are from BGE's Lighting Technical Sheet ({config.get('tech_sheet_edition', '')} edition). "
                      "Products must be on the DLC Qualified Products List, and the final amount is set by BGE at pre-approval.")

    # Savings
    doc.add_heading("Annual savings and payback", level=2)
    pb = f", paying back the {money(n['net'])} net investment in {n['payback_years']:.1f} years" if n["payback_years"] else ""
    doc.add_paragraph(f"The upgrade saves an estimated {money(n['annual_usd'], cents=False)} per year{pb}.")
    s = project["savings"]
    mt = s["maintenance"]
    srows = [["Energy", f"{n['saved_kwh']:,.0f} kWh at {s['rate_cents_per_kwh']:g}¢/kWh ({n['ex_kwh']:,.0f} kWh existing to {n['new_kwh']:,.0f} kWh new)",
              money(n["energy_usd"], cents=False)]]
    if n["maint_usd"]:
        srows.append(["Maintenance", "Fluorescent lamp and ballast replacements avoided", money(n["maint_usd"], cents=False)])
    srows.append(["Total", "", money(n["annual_usd"], cents=False)])
    _table(doc, ["Savings", "Basis", "Per year"], srows, bold_last=1)
    note = (f"Savings assume lights run {s['hours_per_week']:g} hours per week and do not count added sensor savings. "
            f"Energy is valued at {s.get('rate_basis', 'the stated rate')}; your rate on the BGE bill may differ.")
    if n["maint_usd"]:
        note += (f" Maintenance covers fluorescent lamps and ballasts only, at {mt['lamp_minutes']:g} minutes per lamp and "
                 f"{mt['ballast_minutes']:g} minutes per ballast at {money(mt['labor_rate'], cents=False)}/hour plus parts.")
    doc.add_paragraph(note)

    # How it works
    doc.add_heading("How it works", level=2)
    doc.add_paragraph("Nothing is removed or purchased until BGE issues its pre-approval letter.")
    steps = [("Approve.", " Accept this proposal and sign the BGE Terms and Conditions."),
             ("BGE pre-approval.", f" {co['name']} submits the application with the lighting worksheet and product spec sheets."),
             ("Installation.", " New fixtures and sensors are installed and commissioned, and old lamps and ballasts are recycled."),
             ("Close-out.", f" You sign BGE's pre-approval letter confirming the work is complete, and {co['name']} submits the final invoice. "
                            "Pre-approvals for retrofit projects are valid for six months.")]
    for lead, rest in steps:
        _para(doc, [(lead, True), (rest, False)], style="List Number")

    # Acceptance
    doc.add_heading("Acceptance", level=2)
    doc.add_paragraph(f"By signing below, the customer accepts this proposal and authorizes {co['name']} to submit the BGE application.")
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    for row, (a, b) in zip(t.rows, [("Customer", co["name"]), ("Signature:", "Signature:"),
                                     ("Printed name and title:", "Printed name and title:"), ("Date:", "Date:")]):
        row.cells[0].text, row.cells[1].text = a, b
    foot = doc.add_paragraph(f"{co['name']} · {co['street']}, {co['city']}, {co['state']} {co['zip']} · {co['phone']}")
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(out_docx)
    return out_docx, n


def to_pdf(docx_path, out_dir):
    """docx -> pdf with LibreOffice (layout only; never used for BGE math)."""
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
