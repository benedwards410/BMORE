"""The standard KCE audit sheet (.xlsx) and its converter to project.json.

The sheet asks for everything the BGE workbook needs per room, with dropdowns
fed from BGE's own lists (template_tables.json) and the product catalog, so an
audit filled on it converts without guessing. Non-standard audits are handled by
Claude per SKILL.md: copy them onto this sheet first, confirm with the user, then convert.
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .project import load_catalog, load_existing
from .template_map import load_tables

SEP = " · "

PROJECT_FIELDS = [
    ("customer_of_record", "Customer of record (name on the BGE account)", None),
    ("name", "Project name (optional; defaults to customer + address)", None),
    ("site_address", "Site street address", None),
    ("suite", "Suite / unit", None),
    ("city", "City", None),
    ("state", "State", None),
    ("zip", "ZIP", None),
    ("bge_account", "BGE electric account number", None),
    ("building_type", "Building type (BGE list)", "building_types"),
    ("heating_fuel", "Heating fuel", "heating_fuels"),
    ("square_feet", "Building / application square feet", None),
    ("national_account", "National account? (Yes/No)", "yesno"),
    ("business_sector", "Business sector (portal)", None),
    ("business_type", "Business type: Minority-, Woman-, Veteran-owned (comma list, or blank)", None),
    ("contact.name", "Customer contact who signs the T&C: name", None),
    ("contact.title", "Customer contact: title", None),
    ("contact.email", "Customer contact: email", None),
    ("contact.phone", "Customer contact: phone", None),
    ("payee", "BGE incentive paid to (contractor / third_party / customer)", "payee"),
    ("prepared_by", "Prepared by", None),
    ("date", "Audit date (YYYY-MM-DD)", None),
]

AUDIT_COLS = [
    ("id", "Line", 7), ("area", "Area / room", 22), ("int_ext", "Interior / Exterior", 17),
    ("space_type", "Space type (BGE)", 34), ("over_300", "Room over 300 sq ft?", 11), ("space_dim", "Room sq ft", 9),
    ("existing", "Existing fixture", 34), ("existing_qty", "Existing qty", 9), ("existing_controls", "Existing controls in room?", 11),
    ("mount_height_ft", "Mount height ft (high bays)", 11), ("product", "New product", 38), ("qty", "New qty (blank = same)", 10),
    ("ft_per_fixture", "Ft per fixture (strips)", 10), ("control", "Sensor / control", 30), ("control_qty", "Sensor qty (blank = same)", 10),
    ("daylight_qty", "Fixtures in daylight zone", 10), ("ext_zone", "Exterior lighting zone", 30), ("unit_price", "Unit price override", 10),
    ("notes", "Notes", 30),
]
ROWS = 300


def _lists(tables, catalog, existing):
    return {
        "int_ext": tables["int_ext"],
        "space_type": tables["space_types"] + tables["ext_space_types"],
        "yesno": ["Yes", "No"],
        "existing": [f"{k}{SEP}{v['label']}" for k, v in existing["fixtures"].items()],
        "product": [f"{k}{SEP}{v['label']}" for k, v in catalog["fixtures"].items()],
        "control": ["None"] + [f"{k}{SEP}{v['label']}" for k, v in catalog["controls"].items()],
        "ext_zone": tables["ext_zones"],
        "building_types": tables["building_types"],
        "heating_fuels": tables["heating_fuels"],
        "payee": ["contractor", "third_party", "customer"],
    }


def write_template(out_path):
    tables, catalog, existing = load_tables(), load_catalog(), load_existing()
    lists = _lists(tables, catalog, existing)
    wb = Workbook()
    rd = wb.active
    rd.title = "Read me"
    for line in [
        "KCE prescriptive lighting audit (BGE Energy Solutions for Business)",
        "",
        "1. Project tab: fill every row. The T&C contact must be the customer, not KCE.",
        "2. Audit tab: one row per room AND fixture type. Use the dropdowns; they are BGE's own lists.",
        "3. Over 300 sq ft decides whether BGE requires an occupancy sensor in enclosed rooms.",
        "4. Daylight zone: how many fixtures sit near windows. Daylight/DOS sensors (LC 111) are paid only on those.",
        "5. Docks, canopies, garages: say Interior or Exterior on purpose. It changes the measure and the rules.",
        "6. High bays need the mounting height (BGE requires 15 ft or more).",
        "7. Leave New qty / Sensor qty blank for one-for-one.",
        "8. Missing a fixture or product in a dropdown? Type it in Notes and tell Claude; it gets added to the catalog.",
    ]:
        rd.append([line])
    rd["A1"].font = Font(bold=True, size=13)
    rd.column_dimensions["A"].width = 110

    ls = wb.create_sheet("Lists")
    for ci, (key, vals) in enumerate(lists.items(), start=1):
        ls.cell(row=1, column=ci, value=key)
        for ri, v in enumerate(vals, start=2):
            ls.cell(row=ri, column=ci, value=v)
    ls.sheet_state = "hidden"
    ref = {}
    for ci, (key, vals) in enumerate(lists.items(), start=1):
        L = get_column_letter(ci)
        ref[key] = f"=Lists!${L}$2:${L}${len(vals) + 1}"

    pj = wb.create_sheet("Project", 1)
    pj.append(["Field", "Value", "key"])
    head = PatternFill("solid", fgColor="1F3864")
    for c in pj[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    for i, (key, label, lst) in enumerate(PROJECT_FIELDS, start=2):
        pj.append([label, "MD" if key == "state" else ("contractor" if key == "payee" else None), key])
        if lst:
            dv = DataValidation(type="list", formula1=ref[lst], allow_blank=True)
            pj.add_data_validation(dv)
            dv.add(f"B{i}")
    pj.column_dimensions["A"].width = 62
    pj.column_dimensions["B"].width = 44
    pj.column_dimensions["C"].hidden = True

    au = wb.create_sheet("Audit", 2)
    au.append([c[1] for c in AUDIT_COLS])
    au.append([c[0] for c in AUDIT_COLS])
    au.row_dimensions[2].hidden = True
    for ci, (_, _, w) in enumerate(AUDIT_COLS, start=1):
        cell = au.cell(row=1, column=ci)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        au.column_dimensions[get_column_letter(ci)].width = w
    au.row_dimensions[1].height = 45
    au.freeze_panes = "C3"
    for r in range(3, ROWS + 3):
        au.cell(row=r, column=1, value=f"L{r - 2}")
    dv_map = {"int_ext": "int_ext", "space_type": "space_type", "over_300": "yesno", "existing": "existing",
              "existing_controls": "yesno", "product": "product", "control": "control", "ext_zone": "ext_zone"}
    for ci, (key, _, _) in enumerate(AUDIT_COLS, start=1):
        if key in dv_map:
            dv = DataValidation(type="list", formula1=ref[dv_map[key]], allow_blank=True)
            au.add_data_validation(dv)
            L = get_column_letter(ci)
            dv.add(f"{L}3:{L}{ROWS + 2}")
    wb.active = 1
    wb.save(out_path)
    return out_path


def _id(v):
    if not isinstance(v, str):
        return v
    return v.split(SEP)[0].strip()


def _yes(v):
    return isinstance(v, str) and v.strip().lower() in ("yes", "y", "true")


def to_project(path):
    """Filled audit sheet -> (project dict, notes)."""
    wb = load_workbook(path, data_only=True)
    notes = []
    pj, au = wb["Project"], wb["Audit"]
    proj, contact = {}, {}
    for row in pj.iter_rows(min_row=2, values_only=True):
        label, value, key = (row + (None, None, None))[:3]
        if not key or value in (None, ""):
            continue
        if key.startswith("contact."):
            contact[key.split(".", 1)[1]] = value
        else:
            proj[key] = value
    payee = proj.pop("payee", None)
    prepared_by = proj.pop("prepared_by", None)
    date = proj.pop("date", None)
    if "zip" in proj:
        proj["zip"] = str(proj["zip"]).split(".")[0]
    if "bge_account" in proj:
        proj["bge_account"] = str(proj["bge_account"]).split(".")[0]
    if "national_account" in proj:
        proj["national_account"] = _yes(proj["national_account"])
    if isinstance(proj.get("business_type"), str):
        proj["business_type"] = [x.strip() for x in proj["business_type"].split(",") if x.strip()]

    keys = [c.value for c in au[2]]
    lines = []
    for row in au.iter_rows(min_row=3, values_only=True):
        d = dict(zip(keys, row))
        if not d.get("area") and not d.get("existing_qty"):
            continue
        ln = {k: v for k, v in d.items() if v not in (None, "") and k != "notes"}
        for k in ("existing", "product", "control"):
            if k in ln:
                ln[k] = _id(ln[k])
        if ln.get("control") == "None":
            ln.pop("control")
        for k in ("existing_controls",):
            if k in ln:
                ln[k] = _yes(ln[k])
        if d.get("notes"):
            ln["notes"] = d["notes"]
        lines.append(ln)
    project = {"project": proj, "customer_contact": contact, "lines": lines,
               "audit_totals": {"existing_qty": sum(int(l.get("existing_qty") or 0) for l in lines),
                                "lines": len(lines)}}
    if payee:
        project["payee"] = payee
    if prepared_by:
        project["prepared_by"] = prepared_by
    if date:
        project["date"] = str(date)[:10]
    existing, catalog = load_existing(), load_catalog()
    for ln in lines:
        if ln.get("existing") and ln["existing"] not in existing["fixtures"]:
            notes.append(f"{ln['id']}: existing fixture {ln['existing']!r} isn't in the pick list")
        if ln.get("product") and ln["product"] not in catalog["fixtures"]:
            notes.append(f"{ln['id']}: product {ln['product']!r} isn't in the catalog")
    return project, notes
