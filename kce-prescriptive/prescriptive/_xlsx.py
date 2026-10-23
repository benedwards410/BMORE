"""openpyxl loading helpers.

The BGE workbook's stylesheet declares font families above openpyxl's limit of
14, which makes a plain load_workbook() fail. Widen the limit before loading.
Only used for READING; writing the BGE workbook goes through xlsm_writer.
"""

import openpyxl
from openpyxl.styles import fonts
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: F401

fonts.Font.__dict__["family"].max = 255   # raise the limit on the existing descriptor


def load(path, values=True):
    """Read-only workbook. values=True returns Excel's cached results, not formulas."""
    return openpyxl.load_workbook(path, read_only=True, data_only=values, keep_vba=False)


def grid(ws, min_row, max_row, max_col):
    """{'A1': value} for every non-empty cell in the block (read-only friendly)."""
    out = {}
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
        for c in row:
            if c.value is not None and hasattr(c, "coordinate"):
                out[c.coordinate] = c.value
    return out
