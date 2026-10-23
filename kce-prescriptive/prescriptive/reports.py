"""INTERNAL QA workbook and the Excel readback."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import template_map as tm
from ._xlsx import grid, load

FILL = {"BLOCKER": "FEE4E2", "ASK": "FEF0C7", "WARN": "FFFAEB", "INFO": "F2F4F7"}


def qa_workbook(out_xlsx, project, resolved, est, findings, status, readback=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    p = project["project"]
    counts = {lvl: sum(1 for f in findings if f["level"] == lvl) for lvl in ("BLOCKER", "ASK", "WARN", "INFO")}
    rows = [("INTERNAL - not for the customer", ""), ("Project", p.get("name") or p.get("customer_of_record")),
            ("Site", f"{p.get('site_address')}, {p.get('city')} {p.get('zip')}"),
            ("QA status", status)] + [(k, v) for k, v in counts.items()] + [
            ("Estimated lighting incentive", est["lighting"]), ("Estimated control incentive", est["controls"]),
            ("Estimated total incentive", est["total"])]
    if readback:
        rows += [("Excel readback", "matches estimate" if readback.get("ok") else readback.get("message")),
                 ("Excel incentive (L99)", (readback.get("totals") or {}).get("incentive")),
                 ("Excel kWh (L100)", (readback.get("totals") or {}).get("kwh"))]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, color="B42318")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 60

    f = wb.create_sheet("Findings")
    f.append(["Level", "Rule", "Line", "Finding", "Fix"])
    order = {"BLOCKER": 0, "ASK": 1, "WARN": 2, "INFO": 3}
    for x in sorted(findings, key=lambda x: (order[x["level"]], x["rule"])):
        f.append([x["level"], x["rule"], x["line"], x["message"], x["fix"]])
        for c in f[f.max_row]:
            c.fill = PatternFill("solid", fgColor=FILL[x["level"]])
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDE", (10, 6, 7, 90, 50)):
        f.column_dimensions[col].width = w

    ln = wb.create_sheet("Lines")
    head = ["Line", "Area", "Int/Ext", "Space type", "Existing code", "Existing qty", "Product", "Measure", "Qty",
            "Units", "Rate", "Lighting $", "Control", "Control units", "Control $", "Line $", "Why $0 / notes"]
    ln.append(head)
    rb = {x["id"]: x for x in (readback or {}).get("lines", [])}
    if rb:
        ln.cell(row=1, column=len(head) + 1, value="Excel line $")
    for r, e in zip(resolved, est["lines"]):
        row = [r["id"], r.get("area"), r.get("int_ext"), r.get("space_type"), r.get("existing_code"), r.get("existing_qty"),
               (r.get("fixture") or {}).get("model"), r.get("measure"), r.get("qty"), e["incentivized_units"], e["rate"],
               e["lighting"], e["control_code"], e["control_units"], e["control"], e["total"], "; ".join(e["reasons"])]
        if rb:
            row.append(rb.get(r["id"], {}).get("line_incentive"))
        ln.append(row)
    for c in ln[1]:
        c.font = Font(bold=True)
    wb.save(out_xlsx)
    return out_xlsx


def readback(saved_xlsm, resolved, est):
    """Read Excel's results from a workbook that was opened and saved in Excel."""
    wb = load(saved_xlsm, values=True)
    m = grid(wb[tm.MCST], 60, 101, 12)
    totals = {k: m.get(c) for k, c in tm.TOTAL_CELLS.items()}
    if totals["incentive"] is None or totals["kwh"] is None:
        wb.close()
        return {"ok": False, "message": "No calculated totals in this file. Open it in Excel, let it recalculate, save, then run readback again.", "totals": totals}
    n = len(resolved)
    li = grid(wb[tm.LI], tm.LINE_FIRST_ROW, tm.LINE_FIRST_ROW + n - 1, 60)
    wb.close()
    lines, diffs = [], []
    for i, (r, e) in enumerate(zip(resolved, est["lines"])):
        row = tm.LINE_FIRST_ROW + i
        x = {k: li.get(f"{c}{row}") for k, c in tm.LINE_OUTPUTS.items()}
        x["id"] = r["id"]
        lines.append(x)
        def num(v):
            return v if isinstance(v, (int, float)) else 0
        if round(num(x["lighting_incentive"])) != round(e["lighting"]) or round(num(x["control_incentive"])) != round(e["control"]):
            diffs.append({"line": r["id"], "excel_lighting": x["lighting_incentive"], "bot_lighting": e["lighting"],
                          "excel_control": x["control_incentive"], "bot_control": e["control"]})
    ok = not diffs and round(totals["incentive"]) == round(est["total"])
    return {"ok": ok, "totals": totals, "lines": lines, "diffs": diffs,
            "message": "Excel matches the bot's estimate." if ok else
            f"Excel total {totals['incentive']} vs bot {est['total']}; {len(diffs)} line(s) differ."}
